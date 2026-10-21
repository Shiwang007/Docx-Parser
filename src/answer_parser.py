import docx
import re

class AnswerKeyParser:
    def __init__(self, docx_path):
        self.doc = docx.Document(docx_path)
        self.answers = {}

    def parse(self):
        """
        Parses tables or paragraphs to extract question answers.
        For tables: parses tables where even rows are question numbers and odd rows are answers.
        For paragraphs: parses text using pattern matching.
        """
        # Try table parsing first
        for table in self.doc.tables:
            for i in range(0, len(table.rows) - 1, 2):
                num_row = table.rows[i]
                ans_row = table.rows[i+1]
                
                num_cells = [c.text.strip() for c in num_row.cells]
                ans_cells = [c.text.strip() for c in ans_row.cells]
                
                for q_num, q_ans in zip(num_cells, ans_cells):
                    if q_num and q_ans:
                        # Normalize question number (it might be '1' or 'Q1')
                        clean_num = re.sub(r'[^\d]', '', q_num)
                        if clean_num:
                            # Split by comma if multiple correct options exist (e.g., 'B,C')
                            ans_list = [a.strip().upper() for a in q_ans.split(',')]
                            self.answers[clean_num] = ans_list
                            
        # If no answers found in tables, parse paragraphs and table cells
        if not self.answers:
            # Match question numbers followed by dot or parenthesis, e.g. 1. (A) or 1) (B)
            pattern = re.compile(r'(\d+)\s*[\.)]\s*\(([^)]+)\)')
            
            def extract_from_text(text):
                if not text:
                    return
                matches = pattern.findall(text)
                for q_num, ans in matches:
                    ans_list = [a.strip().upper() for a in ans.split(',')]
                    self.answers[q_num] = ans_list

            # Parse paragraphs
            for p in self.doc.paragraphs:
                extract_from_text(p.text)

            # Parse all table cells (useful when text is enclosed in a single table/cell)
            for table in self.doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        extract_from_text(cell.text)

        return self.answers

if __name__ == "__main__":
    import sys
    import json
    if len(sys.argv) > 1:
        parser = AnswerKeyParser(sys.argv[1])
        results = parser.parse()
        print(json.dumps(results, indent=2))
