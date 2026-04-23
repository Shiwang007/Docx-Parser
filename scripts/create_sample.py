import sys
import os
sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), '..\src')))

from docx import Document
from docx.shared import Inches
import os

def create_sample_docx(path):
    doc = Document()
    doc.add_heading('DOCX Parser Test Document', 0)

    # 1. Simple Paragraph
    doc.add_paragraph('This is a simple paragraph to test basic text extraction.')

    # 2. Table with Merged Cells
    doc.add_heading('Table Test', level=1)
    table = doc.add_table(rows=3, cols=3)
    table.style = 'Table Grid'
    
    # Fill table
    for r in range(3):
        for c in range(3):
            table.cell(r, c).text = f"Cell {r},{c}"
            
    # Merge cells: Row 0, Cell 0 and 1
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(0, 0).text = "Merged Header"

    # 3. Image Test (Placeholder text since generating a real image is harder here)
    # Actually, I'll try to add a small dummy image if possible, 
    # but for now let's just use text and formulas.
    doc.add_paragraph('Below this is a formula test (Note: I cannot easily insert OMML via python-docx high-level API, normally this would be an equation).')

    # 4. Mixed Content
    p = doc.add_paragraph('Text before ')
    # In a real Word doc, equations are inserted here.
    
    doc.save(path)
    print(f"[+] Sample document created at {path}")

if __name__ == "__main__":
    if not os.path.exists('test_data'):
        os.makedirs('test_data')
    create_sample_docx('test_data/sample.docx')
