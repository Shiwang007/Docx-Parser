import sys
import os
sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), '..\src')))


from docx import Document

doc = Document('DPP-01_Matrices_Maths_Nishant Sir.docx')
with open('docx_content.txt', 'w', encoding='utf-8') as f:
    for i, p in enumerate(doc.paragraphs):
        f.write(f"{i}: {p.text}\n")
print(f"Total paragraphs: {len(doc.paragraphs)}")
