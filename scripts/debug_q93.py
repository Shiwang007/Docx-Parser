import sys
import os
sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), '..\src')))

from docx import Document
import os

def debug_q93_v2():
    doc = Document('Test-03_Class-XI+XII_AITS_NEET 2026.docx')
    found = False
    for i, p in enumerate(doc.paragraphs):
        if 'Housefly belongs to the family' in p.text:
            print(f"Found Q93 at index {i}")
            for j in range(i, i+5):
                if j < len(doc.paragraphs):
                    text = doc.paragraphs[j].text
                    print(f"P[{j}] text: {repr(text)}")
            found = True
            break
    if not found:
        print("Could not find Q93")

if __name__ == "__main__":
    debug_q93_v2()
